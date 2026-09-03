"""Convert the supplied Riven good-roll workbook into runtime JSON data."""

from __future__ import annotations

import argparse
import json
import time
from pathlib import Path

from riven_roll_rules import parse_negative_expression, parse_rule_expression


SHEETS = {
    "primary": (1, 2, 6, 9),
    "secondary": (1, 2, 6, 9),
    "melee": (1, 2, 7, 10),
    "archgun": (1, 2, 8, 10),
    "robotic": (1, 2, 5, 7),
}


def key(value: object) -> str:
    import re

    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


def import_workbook(source: Path) -> dict:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("This one-time importer requires openpyxl") from exc

    workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    missing = sorted(set(SHEETS) - set(workbook.sheetnames))
    if missing:
        raise ValueError(f"Workbook is missing sheets: {', '.join(missing)}")

    rules: list[dict] = []
    for sheet_name, (weapon_col, positive_col, negative_col, notes_col) in SHEETS.items():
        sheet = workbook[sheet_name]
        for row_number, source_row in enumerate(sheet.iter_rows(values_only=True), 1):
            row = tuple(source_row) + (None,) * 12
            weapon = str(row[weapon_col - 1] or "").strip()
            expression = str(row[positive_col - 1] or "").strip()
            if not weapon or not expression or weapon.casefold().startswith(("weapon", "good")):
                continue
            negative_expression = str(row[negative_col - 1] or "").strip()
            notes = str(row[notes_col - 1] or "").strip()
            rules.append({
                "weapon": weapon,
                "key": key(weapon),
                "category": sheet_name,
                "positive_expression": expression,
                "negative_expression": negative_expression,
                "harmless_negatives": list(parse_negative_expression(negative_expression)),
                "alternatives": list(parse_rule_expression(expression)),
                "notes": notes,
                "source_cell": f"{sheet_name}!A{row_number}",
            })

    # Hybrid weapons can intentionally appear on more than one tab. The live
    # marketplace names the second family with its mode (for example,
    # ``Vinquibus (Melee)``), so retain both instead of overwriting either.
    merged: dict[str, dict] = {}
    for rule in rules:
        existing = merged.get(rule["key"])
        if existing is not None and rule["category"] != existing["categories"][0]:
            rule["weapon"] = f"{rule['weapon']} ({rule['category'].title()})"
            rule["key"] = key(rule["weapon"])
            existing = merged.get(rule["key"])
        if existing is None:
            rule["categories"] = [rule.pop("category")]
            rule["source_cells"] = [rule.pop("source_cell")]
            merged[rule["key"]] = rule
            continue
        existing["categories"].append(rule["category"])
        existing["source_cells"].append(rule["source_cell"])
        existing["positive_expression"] += f" | {rule['category']}: {rule['positive_expression']}"
        existing["negative_expression"] += f" | {rule['category']}: {rule['negative_expression']}"
        existing["harmless_negatives"] = sorted(set(existing["harmless_negatives"]) | set(rule["harmless_negatives"]))
        existing["alternatives"].extend(rule["alternatives"])
        if rule["notes"]:
            existing["notes"] = " | ".join(value for value in (existing["notes"], rule["notes"]) if value)
    rules = list(merged.values())
    return {
        "source": source.name,
        "generated_at": time.time(),
        "rule_count": len(rules),
        "requirements": {
            "minimum_desired_positives": 2,
            "harmless_negative_required": True,
            "dead_third_positive_allowed": False,
        },
        "rows": rules,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Import curated Riven resale-roll rules from XLSX")
    parser.add_argument("source", type=Path)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).resolve().parent / "data" / "rivens" / "roll_rules.json",
    )
    args = parser.parse_args()
    payload = import_workbook(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Imported {payload['rule_count']} weapon roll rules to {args.output}")


if __name__ == "__main__":
    main()
